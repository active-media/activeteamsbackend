from __future__ import annotations

import asyncio
import csv
import io
import logging
import os
import tempfile
from datetime import date
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from Reports.twelve_tasks import sb_get_twelve_tasks_report, export_twelve_tasks_xlsx
from supabase_helpers.excel_export import build_cells_excel
from supabase_helpers.school_cell_report import sb_get_school_cell_report, build_school_cell_excel
from supabase_helpers.service_targets import sb_get_service_target_report
from supabase_helpers.supabase_connection import supabase_admin
from supabase_helpers.supabase_stats import sb_get_reporting_dashboard

from .envelope import (
    cells_graph_envelope,
    dashboard_envelope,
    school_cell_envelope,
    service_target_envelope,
    twelve_tasks_envelope,
)
from .periods import ResolvedPeriod, period_to_week_ids, resolve_period
from .schemas import (
    DashboardEnvelope,
    ReportCatalogItem,
    ReportErrorResponse,
    ReportEnvelope,
    ReportNotImplementedResponse,
    ReportRunRecord,
)
from .scope import ReportScope, get_report_scope


logger = logging.getLogger(__name__)
router = APIRouter(prefix="/reporting", tags=["Reporting API"])

PERIOD_TOKENS = [
    "today", "thisWeek", "thisMonth", "previous7", "previousWeek",
    "previousMonth", "thisQuarter", "yearly", "YYYY-MM",
]
REPORT_CATALOG = [
    ReportCatalogItem(
        report_type="twelve_tasks", title="Twelve Tasks",
        description="Completed task counts for tracked people and targets.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}],
        supports_export=["csv", "xlsx"], implemented=True,
    ),
    ReportCatalogItem(
        report_type="cells_graph", title="Cells Graph",
        description="Cell attendance and active-cell trends.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}, {"name": "campus"}],
        supports_export=["csv", "xlsx"], implemented=True,
    ),
    ReportCatalogItem(
        report_type="life_class", title="Life Class",
        description="Life Class attendance report.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}, {"name": "campus"}],
        supports_export=["csv", "xlsx"], implemented=False,
    ),
    ReportCatalogItem(
        report_type="school_of_leaders", title="School of Leaders",
        description="School of Leaders attendance report.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}, {"name": "campus"}],
        supports_export=["csv", "xlsx"], implemented=False,
    ),
    ReportCatalogItem(
        report_type="plan_40", title="Plan 40",
        description="Plan 40 progress report.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}, {"name": "campus"}],
        supports_export=["csv", "xlsx"], implemented=False,
    ),
    ReportCatalogItem(
        report_type="staff_interns_youth", title="Staff/Interns/Youth",
        description="Staff, interns, and youth report.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}, {"name": "campus"}],
        supports_export=["csv", "xlsx"], implemented=False,
    ),
    ReportCatalogItem(
        report_type="school_cell", title="School Cell",
        description="Weekly School Cell attendance and leader summary.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}],
        supports_export=["csv", "xlsx"], implemented=True,
    ),
    ReportCatalogItem(
        report_type="service_target", title="Service Target",
        description="Weekly leader service targets.",
        params=[{"name": "period", "allowed_values": PERIOD_TOKENS}, {"name": "start_date"}, {"name": "end_date"}],
        supports_export=["csv", "xlsx"], implemented=True,
    ),
]
CATALOG_BY_TYPE = {item.report_type: item for item in REPORT_CATALOG}
LEGACY_SCHEDULE_TYPES = {
    "overall_church_performance": "cells_graph",
    "cells_report": "cells_graph",
    "life_class_report": "life_class",
    "school_of_leaders_report": "school_of_leaders",
    "plan_40_report": "plan_40",
    "school_cell_report": "school_cell",
    "service_target_report": "service_target",
    "twelve_tasks_report": "twelve_tasks",
    "staff_interns_youth_report": "staff_interns_youth",
}
NOT_IMPLEMENTED_RESPONSE = {
    501: {
        "model": ReportNotImplementedResponse,
        "description": "The report or requested export is not implemented.",
    }
}
COMMON_ERROR_RESPONSES = {
    status_code: {"model": ReportErrorResponse, "description": description}
    for status_code, description in {
        401: "Authentication is required.",
        403: "The user role or organization scope is not permitted.",
        422: "The request parameters are invalid.",
        500: "The report could not be calculated or retrieved.",
    }.items()
}
REPORT_ERROR_RESPONSES = {**COMMON_ERROR_RESPONSES, **NOT_IMPLEMENTED_RESPONSE}
REPORT_ERROR_RESPONSES[404] = {
    "model": ReportErrorResponse,
    "description": "The requested report type or report run was not found.",
}
EXPORT_RESPONSES = {
    **REPORT_ERROR_RESPONSES,
    200: {
        "description": "CSV or XLSX report download.",
        "content": {
            "text/csv": {"schema": {"type": "string"}},
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {
                "schema": {"type": "string", "format": "binary"}
            },
        },
    },
}


def _not_implemented(report_type: str, message: str) -> HTTPException:
    return HTTPException(
        status_code=501,
        detail={"code": "not_implemented", "report_type": report_type, "message": message},
    )


def _resolved_period(period: str | None, start_date: str | None, end_date: str | None) -> ResolvedPeriod:
    return resolve_period(period, start_date, end_date)


def _dates(period: ResolvedPeriod) -> tuple[str, str]:
    return period.start.isoformat(), period.end.isoformat()


async def _dashboard_data(scope: ReportScope, period: ResolvedPeriod, campus: str | None) -> dict:
    start, end = _dates(period)
    return await asyncio.to_thread(
        sb_get_reporting_dashboard,
        period.token,
        scope.org_filter,
        campus,
        scope.visible_leader_ids,
        date.fromisoformat(start),
        date.fromisoformat(end),
    )


async def _service_target_data(scope: ReportScope, period: ResolvedPeriod) -> dict:
    rows = []
    for week_id in period_to_week_ids(period):
        result = await asyncio.to_thread(
            sb_get_service_target_report,
            week_id,
            scope.org_filter,
            scope.visible_leader_ids,
        )
        rows.extend(result.get("targets") or [])
    return {
        "targets": rows,
        "total_target": sum(row.get("target_count") or 0 for row in rows),
        "total_targets": len(rows),
    }


async def _twelve_tasks_data(
    scope: ReportScope,
    period: ResolvedPeriod,
    nullable_targets: bool = True,
) -> dict:
    return await asyncio.to_thread(
        sb_get_twelve_tasks_report,
        period.token,
        scope.org_filter,
        None,
        period.start,
        period.end,
        scope.visible_leader_ids,
        nullable_targets,
    )


async def _optional_dashboard_data(section: str, data_awaitable):
    try:
        return await data_awaitable
    except Exception:
        logger.exception("Unable to calculate reporting dashboard %s data", section)
        return None


async def _report_envelope(
    report_type: str,
    scope: ReportScope,
    period: ResolvedPeriod,
    campus: str | None,
) -> ReportEnvelope:
    if report_type not in CATALOG_BY_TYPE:
        raise HTTPException(status_code=404, detail=f"Unknown report_type '{report_type}'")
    if not CATALOG_BY_TYPE[report_type].implemented:
        raise _not_implemented(report_type, f"The {report_type} report is not implemented.")

    try:
        if report_type == "twelve_tasks":
            return twelve_tasks_envelope(await _twelve_tasks_data(scope, period), period)
        if report_type == "school_cell":
            start, end = _dates(period)
            data = await asyncio.to_thread(
                sb_get_school_cell_report, scope.org_filter, start, end, scope.visible_leader_ids
            )
            return school_cell_envelope(data, period)
        if report_type == "service_target":
            return service_target_envelope(await _service_target_data(scope, period), period)
        data = await _dashboard_data(scope, period, campus)
        return cells_graph_envelope(data, period, campus)
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Unable to calculate report %s", report_type)
        raise HTTPException(status_code=500, detail=f"Unable to calculate report '{report_type}'") from exc


@router.get(
    "/catalog",
    response_model=list[ReportCatalogItem],
    responses=COMMON_ERROR_RESPONSES,
)
async def get_report_catalog(scope: ReportScope = Depends(get_report_scope)):
    return REPORT_CATALOG


@router.get("/dashboard", response_model=DashboardEnvelope, responses=COMMON_ERROR_RESPONSES)
async def get_reporting_dashboard(
    period: str | None = Query(None),
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    campus: str | None = Query(None, min_length=1, max_length=120),
    scope: ReportScope = Depends(get_report_scope),
):
    resolved = _resolved_period(period, start_date, end_date)
    try:
        data, service_data, twelve_data = await asyncio.gather(
            _dashboard_data(scope, resolved, campus),
            _optional_dashboard_data("Service Target", _service_target_data(scope, resolved)),
            _optional_dashboard_data("Twelve Tasks", _twelve_tasks_data(scope, resolved)),
        )
        return dashboard_envelope(data, resolved, campus, service_data, twelve_data)
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Unable to calculate reporting dashboard")
        raise HTTPException(status_code=500, detail="Unable to calculate reporting dashboard") from exc


@router.get("/reports/{report_type}", response_model=ReportEnvelope, responses=REPORT_ERROR_RESPONSES)
async def get_report(
    report_type: str,
    period: str | None = Query(None),
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    campus: str | None = Query(None, min_length=1, max_length=120),
    scope: ReportScope = Depends(get_report_scope),
):
    resolved = _resolved_period(period, start_date, end_date)
    return await _report_envelope(report_type, scope, resolved, campus)


def _csv_content(envelope: ReportEnvelope) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    for index, table in enumerate(envelope.tables):
        if index:
            writer.writerow([])
        writer.writerow([table.title])
        writer.writerow([column.label for column in table.columns])
        for row in table.rows:
            writer.writerow([row.get(column.key) for column in table.columns])
    return output.getvalue()


def _generic_xlsx(envelope: ReportEnvelope) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in envelope.tables:
        sheet = workbook.create_sheet(title=table.key[:31] or "Report")
        sheet.append([table.title])
        sheet.append([column.label for column in table.columns])
        for row in table.rows:
            sheet.append([row.get(column.key) for column in table.columns])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


async def _xlsx_for_report(
    report_type: str,
    scope: ReportScope,
    period: ResolvedPeriod,
    campus: str | None,
    envelope: ReportEnvelope,
) -> bytes:
    start, end = _dates(period)
    if report_type == "cells_graph":
        return await asyncio.to_thread(
            build_cells_excel, scope.org_filter, start, end, scope.visible_leader_ids, campus
        )
    if report_type == "school_cell":
        return await asyncio.to_thread(
            build_school_cell_excel, scope.org_filter, start, end, scope.visible_leader_ids
        )
    if report_type == "twelve_tasks":
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        path = handle.name
        handle.close()
        try:
            report = await _twelve_tasks_data(scope, period, nullable_targets=False)
            await asyncio.to_thread(export_twelve_tasks_xlsx, report, path)
            return await asyncio.to_thread(lambda: open(path, "rb").read())
        finally:
            if os.path.exists(path):
                os.unlink(path)
    return await asyncio.to_thread(_generic_xlsx, envelope)


@router.get("/reports/{report_type}/export", response_class=Response, responses=EXPORT_RESPONSES)
async def export_report(
    report_type: str,
    format: str = Query(..., pattern="^(csv|xlsx|pdf)$"),
    period: str | None = Query(None),
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    campus: str | None = Query(None, min_length=1, max_length=120),
    scope: ReportScope = Depends(get_report_scope),
):
    resolved = _resolved_period(period, start_date, end_date)
    if report_type not in CATALOG_BY_TYPE:
        raise HTTPException(status_code=404, detail=f"Unknown report_type '{report_type}'")
    if not CATALOG_BY_TYPE[report_type].implemented:
        raise _not_implemented(report_type, f"The {report_type} report is not implemented.")
    if format == "pdf":
        raise _not_implemented(report_type, "PDF export is not implemented.")
    envelope = await _report_envelope(report_type, scope, resolved, campus)
    filename = f"{report_type}_{resolved.start.isoformat()}_{resolved.end.isoformat()}.{format}"
    try:
        if format == "csv":
            return Response(
                content=_csv_content(envelope), media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'},
            )
        content = await _xlsx_for_report(report_type, scope, resolved, campus, envelope)
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Unable to export report %s", report_type)
        raise HTTPException(status_code=500, detail=f"Unable to export report '{report_type}'") from exc


def _history_organization(scope: ReportScope, requested_organization: str | None) -> str:
    own_organization = (scope.org_filter or {}).get("Organization")
    if scope.is_super:
        if not requested_organization:
            raise HTTPException(status_code=422, detail="organization is required for super-admin history access")
        return requested_organization
    if requested_organization and requested_organization != own_organization:
        raise HTTPException(status_code=403, detail="Cannot access report history for another organization")
    if not own_organization:
        raise HTTPException(status_code=403, detail="Organization not associated with user")
    return own_organization


def _history_query(scope: ReportScope, organization: str):
    query = supabase_admin.table("report_runs").select("*")
    return query.eq("org", organization)


@router.post(
    "/reports/{report_type}/generate",
    response_model=ReportRunRecord,
    responses=REPORT_ERROR_RESPONSES,
)
async def generate_report(
    report_type: str,
    period: str | None = Query(None),
    start_date: str | None = Query(None),
    end_date: str | None = Query(None),
    campus: str | None = Query(None, min_length=1, max_length=120),
    organization: str | None = Query(None, min_length=1),
    scope: ReportScope = Depends(get_report_scope),
):
    resolved = _resolved_period(period, start_date, end_date)
    envelope = await _report_envelope(report_type, scope, resolved, campus)
    payload = envelope.model_dump(mode="json")
    params = {"period": resolved.token, "start_date": resolved.start.isoformat(), "end_date": resolved.end.isoformat(), "campus": campus}
    org = _history_organization(scope, organization)
    run = {
        "org": org,
        "report_type": report_type,
        "params": params,
        "payload": payload,
        "generated_by": scope.user_id,
    }
    try:
        result = await asyncio.to_thread(
            lambda: supabase_admin.table("report_runs").insert(run).execute()
        )
    except Exception as exc:
        logger.exception("Unable to persist report run")
        raise HTTPException(status_code=500, detail="Unable to persist report run; apply the reporting migration") from exc
    rows = result.data or []
    if not rows:
        raise HTTPException(status_code=500, detail="Report run was not returned after persistence")
    return rows[0]


@router.get("/history", response_model=list[ReportRunRecord], responses=COMMON_ERROR_RESPONSES)
async def get_reporting_history(
    limit: int = Query(50, ge=1, le=200),
    organization: str | None = Query(None, min_length=1),
    scope: ReportScope = Depends(get_report_scope),
):
    try:
        org = _history_organization(scope, organization)
        result = await asyncio.to_thread(lambda: _history_query(scope, org).order("created_at", desc=True).limit(limit).execute())
        return result.data or []
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Unable to load report history")
        raise HTTPException(status_code=500, detail="Unable to load report history") from exc


@router.get("/history/{run_id}", response_model=ReportRunRecord, responses=REPORT_ERROR_RESPONSES)
async def get_reporting_history_item(
    run_id: str,
    organization: str | None = Query(None, min_length=1),
    scope: ReportScope = Depends(get_report_scope),
):
    try:
        org = _history_organization(scope, organization)
        query = _history_query(scope, org).eq("id", run_id).limit(1)
        result = await asyncio.to_thread(lambda: query.execute())
        row = (result.data or [None])[0]
        if not row:
            raise HTTPException(status_code=404, detail="Report run not found")
        return row
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Unable to load report run %s", run_id)
        raise HTTPException(status_code=500, detail="Unable to load report run") from exc