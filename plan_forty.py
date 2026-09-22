"""
plan_forty.py
=============

Exports "Plan 40" cell/discipleship attendance for a term into an Excel
workbook shaped like ``Plan_40_Term_2_2026.xlsx``:

    - "MEN"   sheet: one row per male attendee, grouped under the leader
      of the cell they attend most, with an "X" in each week column they
      were checked in, a running "Total Attendance" count, and a row for
      each leader themself (X = their cell met / they didn't "did not meet").
    - "WOMEN" sheet: same shape, for female attendees.
    - "STATS W1" .. "STATS W5": one sheet per week, listing every leader
      (MEN section then WOMEN section) with the head-count captured for
      their session that week, plus subtotal and grand-total formulas.

HOW THIS MAPS ONTO THE SCHEMA
------------------------------
- ``event_types.name`` identifies the program (default: "Plan 40").
- Each row in ``events`` is treated as one leader's cell/home-group.
  ``event_leader`` becomes the "Leader at 12" for everyone who attends
  that event.
- Each row in ``event_sessions`` is one weekly meeting of an event.
  ``week_identifier`` is matched against WEEK_LABELS below (with a
  couple of fallbacks - see ``week_label_for_session``), and
  ``checked_in_count`` feeds the STATS sheets. ``is_did_not_meet`` marks
  a week the cell didn't run (no X for the leader that week).
- Each row in ``event_session_attendees`` links a person to a session;
  ``is_checked_in`` marks real attendance.

TWO THINGS THE SUPPLIED SCHEMA DOESN'T CONTAIN, WHICH YOU MAY NEED TO WIRE UP
------------------------------------------------------------------------------
1. Gender (which sheet a person/leader belongs to). There's no gender
   column on ``events``/``event_session_attendees`` in the schema given,
   so it's inferred from the event name (see ``classify_gender``). If
   your event names don't contain "men"/"women" etc., either rename them
   consistently or fill in EVENT_GENDER_OVERRIDES below with
   {event_id: "MEN"/"WOMEN"}.
2. "Leader at 144" (the leader's own leader). Nothing in the schema
   models this hierarchy, so it's a manual lookup in LEADER_AT_144
   below. Swap it for a real query if you have a leadership table.

Run with e.g.:
    python plan_forty.py --start 2026-04-20 --end 2026-05-24 \
        --organization "My Church" --output Plan_40_Term_2_2026.xlsx
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from supabase_helpers.supabase_connection import supabase

# --------------------------------------------------------------------------- #
# CONFIG - edit these to match your term / data
# --------------------------------------------------------------------------- #

EVENT_TYPE_NAME = "Plan 40"           # matches event_types.name
ORGANIZATION: Optional[str] = None    # e.g. "My Church" to filter events."Organization"; None = all

TERM_LABEL = "Term 2 2026"
TERM_START = date(2026, 4, 20)        # <-- adjust to the real term start
TERM_END = date(2026, 5, 24)          # <-- adjust to the real term end

WEEK_LABELS = ["Week 1", "Week 2", "Week 3", "Week 4", "Week 5"]

# event_id -> "MEN"/"WOMEN", for events whose name doesn't say which.
EVENT_GENDER_OVERRIDES: dict[str, str] = {
    # "3f1e5c2a-...": "WOMEN",
}

# leader_at_12 name -> leader_at_144 name (no such table in the schema given).
LEADER_AT_144: dict[str, str] = {
    # "Kayla Enslin": "Tracy Bebel",
}

MEN_KEYWORDS = ("men", "man", "guys", "brother")
WOMEN_KEYWORDS = ("women", "woman", "ladies", "lady", "sister")

DEFAULT_OUTPUT_PATH = f"Plan_40_{TERM_LABEL.replace(' ', '_')}_export.xlsx"

# --------------------------------------------------------------------------- #
# Styling helpers - approximates the look of Plan_40_Term_2_2026.xlsx
# --------------------------------------------------------------------------- #

THIN = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Arial", bold=True)
BODY_FONT = Font(name="Arial", bold=False)
CENTER = Alignment(horizontal="center", vertical="bottom")
LEFT = Alignment(horizontal="general", vertical="bottom")


def _style_header(cell) -> None:
    cell.font = HEADER_FONT
    cell.border = BORDER
    cell.alignment = CENTER


def _style_body(cell, centered: bool = False) -> None:
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = CENTER if centered else LEFT


# --------------------------------------------------------------------------- #
# Data access
# --------------------------------------------------------------------------- #

def fetch_event_type_id(name: str) -> str:
    resp = (
        supabase.table("event_types")
        .select("event_type_id, name")
        .ilike("name", name)
        .limit(1)
        .execute()
    )
    rows = resp.data or []
    if not rows:
        raise RuntimeError(f"No event_types row found with name like {name!r}")
    return rows[0]["event_type_id"]


def fetch_events(event_type_id: str, organization: Optional[str]) -> list[dict]:
    query = (
        supabase.table("events")
        .select(
            'event_id, event_name, event_leader, event_leader_email, '
            'is_active, "Organization"'
        )
        .eq("event_type_id", event_type_id)
    )
    if organization:
        query = query.eq("Organization", organization)
    return query.execute().data or []


def classify_gender(event: dict) -> Optional[str]:
    override = EVENT_GENDER_OVERRIDES.get(event["event_id"])
    if override:
        return override
    name = (event.get("event_name") or "").lower()
    if any(k in name for k in WOMEN_KEYWORDS):
        return "WOMEN"
    if any(k in name for k in MEN_KEYWORDS):
        return "MEN"
    return None


def fetch_sessions(event_ids: list[str], start: date, end: date) -> list[dict]:
    if not event_ids:
        return []
    resp = (
        supabase.table("event_sessions")
        .select(
            "session_id, event_id, session_date, week_identifier, "
            "is_did_not_meet, checked_in_count, total_headcounts"
        )
        .in_("event_id", event_ids)
        .gte("session_date", start.isoformat())
        .lte("session_date", end.isoformat())
        .execute()
    )
    return resp.data or []


def fetch_attendees(session_ids: list[str]) -> list[dict]:
    if not session_ids:
        return []
    resp = (
        supabase.table("event_session_attendees")
        .select(
            "session_id, event_id, mongo_person_id, full_name, email, "
            "is_checked_in"
        )
        .in_("session_id", session_ids)
        .execute()
    )
    return resp.data or []


# --------------------------------------------------------------------------- #
# Aggregation
# --------------------------------------------------------------------------- #

def week_label_for_session(session: dict, term_start: date) -> Optional[str]:
    """Map a session onto one of WEEK_LABELS, tolerating a few identifier styles."""
    wi = (session.get("week_identifier") or "").strip()
    if wi in WEEK_LABELS:
        return wi

    digits = "".join(ch for ch in wi if ch.isdigit())
    if digits and 1 <= int(digits) <= len(WEEK_LABELS):
        return WEEK_LABELS[int(digits) - 1]

    session_date = session.get("session_date")
    if session_date:
        d = date.fromisoformat(str(session_date)[:10])
        week_index = (d - term_start).days // 7
        if 0 <= week_index < len(WEEK_LABELS):
            return WEEK_LABELS[week_index]

    return None


def build_report(term_start: date, term_end: date, organization: Optional[str]):
    """Returns (people, leaders) dicts used to populate every sheet."""
    event_type_id = fetch_event_type_id(EVENT_TYPE_NAME)
    events = fetch_events(event_type_id, organization)
    events_by_id = {e["event_id"]: e for e in events}

    gender_by_event: dict[str, str] = {}
    for event in events:
        gender = classify_gender(event)
        if gender is None:
            print(
                f"[warn] could not classify gender for event "
                f"{event.get('event_name')!r} ({event['event_id']}) - skipping. "
                f"Add it to EVENT_GENDER_OVERRIDES to include it."
            )
            continue
        gender_by_event[event["event_id"]] = gender

    event_ids = list(gender_by_event)
    sessions = fetch_sessions(event_ids, term_start, term_end)
    sessions_by_id = {s["session_id"]: s for s in sessions}
    attendees = fetch_attendees(list(sessions_by_id))

    people: dict[str, dict] = {}
    leaders: dict[str, dict] = {}

    # Leader rows: did their cell meet each week, and what was the headcount.
    for session in sessions:
        event = events_by_id.get(session["event_id"])
        gender = gender_by_event.get(session["event_id"])
        if not event or not gender:
            continue
        label = week_label_for_session(session, term_start)
        if not label:
            continue

        leader_name = event.get("event_leader") or "Unassigned"
        leader = leaders.setdefault(
            leader_name,
            {
                "gender": gender,
                "weeks": {w: False for w in WEEK_LABELS},
                "checked_in_by_week": {w: 0 for w in WEEK_LABELS},
            },
        )
        if not session.get("is_did_not_meet"):
            leader["weeks"][label] = True
        leader["checked_in_by_week"][label] += session.get("checked_in_count") or 0

    # Attendee rows: which weeks they were checked in, and which event they
    # attend most (used to pick their "Leader at 12").
    for attendee in attendees:
        if not attendee.get("is_checked_in"):
            continue
        session = sessions_by_id.get(attendee["session_id"])
        if not session:
            continue
        event = events_by_id.get(session["event_id"])
        gender = gender_by_event.get(session["event_id"])
        if not event or not gender:
            continue
        label = week_label_for_session(session, term_start)
        if not label:
            continue

        key = attendee.get("mongo_person_id") or (
            f"{(attendee.get('full_name') or '').strip().lower()}|"
            f"{(attendee.get('email') or '').strip().lower()}"
        )
        person = people.setdefault(
            key,
            {
                "name": attendee.get("full_name") or "(unnamed)",
                "gender": gender,
                "weeks": {w: False for w in WEEK_LABELS},
                "event_counts": defaultdict(int),
            },
        )
        person["weeks"][label] = True
        person["event_counts"][event["event_id"]] += 1

    for person in people.values():
        if not person["event_counts"]:
            person["leader_at_12"] = ""
            person["leader_at_144"] = ""
            continue
        home_event_id = max(person["event_counts"], key=person["event_counts"].get)
        home_event = events_by_id.get(home_event_id, {})
        leader_at_12 = home_event.get("event_leader") or ""
        person["leader_at_12"] = leader_at_12
        person["leader_at_144"] = LEADER_AT_144.get(leader_at_12, "")

    return people, leaders


# --------------------------------------------------------------------------- #
# Sheet writers
# --------------------------------------------------------------------------- #

def _write_week_cells(ws, row: int, first_col: int, weeks: dict[str, bool]) -> None:
    for i, label in enumerate(WEEK_LABELS):
        cell = ws.cell(row=row, column=first_col + i, value="X" if weeks[label] else None)
        _style_body(cell, centered=True)


def write_attendance_sheet(wb: Workbook, gender: str, people: dict, leaders: dict) -> None:
    ws = wb.create_sheet(gender)
    ws.column_dimensions["A"].width = 19.5
    ws.column_dimensions["B"].width = 20.25

    first_week_col = 4  # column D
    last_week_col = first_week_col + len(WEEK_LABELS) - 1
    total_col = last_week_col + 1
    first_week_letter = get_column_letter(first_week_col)
    last_week_letter = get_column_letter(last_week_col)

    headers = ["Name & Surname", "Leader at 12", "Leader at 144", *WEEK_LABELS, "Total Attendance"]
    for col, text in enumerate(headers, start=1):
        _style_header(ws.cell(row=1, column=col, value=text))

    gender_people = [p for p in people.values() if p["gender"] == gender]
    gender_leaders = {n: l for n, l in leaders.items() if l["gender"] == gender}

    grouped: dict[str, list[dict]] = defaultdict(list)
    for person in gender_people:
        grouped[person.get("leader_at_12", "")].append(person)

    row = 2
    for leader_name in sorted(grouped):
        for person in sorted(grouped[leader_name], key=lambda p: p["name"]):
            ws.cell(row=row, column=1, value=person["name"])
            ws.cell(row=row, column=2, value=person.get("leader_at_12", ""))
            ws.cell(row=row, column=3, value=person.get("leader_at_144", ""))
            for col in (1, 2, 3):
                _style_body(ws.cell(row=row, column=col))
            _write_week_cells(ws, row, first_week_col, person["weeks"])
            total_cell = ws.cell(
                row=row, column=total_col,
                value=f"=COUNTA({first_week_letter}{row}:{last_week_letter}{row})",
            )
            _style_header(total_cell)
            row += 1

    for leader_name in sorted(gender_leaders):
        leader = gender_leaders[leader_name]
        _style_body(ws.cell(row=row, column=1, value=leader_name))
        _style_body(ws.cell(row=row, column=2, value=""))
        _style_body(ws.cell(row=row, column=3, value=""))
        _write_week_cells(ws, row, first_week_col, leader["weeks"])
        total_cell = ws.cell(
            row=row, column=total_col,
            value=f"=COUNTA({first_week_letter}{row}:{last_week_letter}{row})",
        )
        _style_header(total_cell)
        row += 1

    last_data_row = row - 1
    for i in range(len(WEEK_LABELS)):
        col = first_week_col + i
        col_letter = get_column_letter(col)
        footer_cell = ws.cell(
            row=row, column=col,
            value=f"=COUNTA({col_letter}2:{col_letter}{last_data_row})",
        )
        _style_header(footer_cell)


def write_stats_sheet(wb: Workbook, week_index: int, week_label: str, leaders: dict) -> None:
    ws = wb.create_sheet(f"STATS W{week_index + 1}")
    ws.cell(row=1, column=1, value=week_label.upper()).font = HEADER_FONT

    def section(start_row: int, gender: str) -> int:
        ws.cell(row=start_row, column=1, value=gender).font = HEADER_FONT
        ws.cell(row=start_row + 1, column=1, value="Name & Surname").font = HEADER_FONT
        ws.cell(row=start_row + 1, column=2, value="Number ").font = HEADER_FONT

        names = sorted(n for n, l in leaders.items() if l["gender"] == gender)
        r = start_row + 2
        for name in names:
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=leaders[name]["checked_in_by_week"].get(week_label, 0))
            r += 1

        subtotal_row = r
        first_data_row = start_row + 2
        if r > first_data_row:
            ws.cell(row=subtotal_row, column=3, value=f"=SUM(B{first_data_row}:B{subtotal_row - 1})")
        else:
            ws.cell(row=subtotal_row, column=3, value=0)
        return subtotal_row

    men_subtotal_row = section(2, "MEN")
    women_subtotal_row = section(men_subtotal_row + 2, "WOMEN")
    grand_row = women_subtotal_row + 1
    ws.cell(row=grand_row, column=3, value=f"=SUM(C{men_subtotal_row},C{women_subtotal_row})")


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Plan 40 attendance to xlsx")
    parser.add_argument("--start", type=date.fromisoformat, default=TERM_START,
                         help="Term start date, YYYY-MM-DD")
    parser.add_argument("--end", type=date.fromisoformat, default=TERM_END,
                         help="Term end date, YYYY-MM-DD")
    parser.add_argument("--organization", default=ORGANIZATION,
                         help='Value to filter events."Organization" by (default: no filter)')
    parser.add_argument("--output", default=DEFAULT_OUTPUT_PATH,
                         help="Path to write the .xlsx file to")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    people, leaders = build_report(args.start, args.end, args.organization)

    wb = Workbook()
    wb.remove(wb.active)

    write_attendance_sheet(wb, "MEN", people, leaders)
    write_attendance_sheet(wb, "WOMEN", people, leaders)
    for i, label in enumerate(WEEK_LABELS):
        write_stats_sheet(wb, i, label, leaders)

    wb.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()