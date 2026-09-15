import openpyxl
wb = openpyxl.load_workbook('twelve_tasks_report.xlsx')
ws = wb.active
print('Sheet:', ws.title)
for r in range(1, ws.max_row + 1):
    row_vals = []
    for c in range(1, 6):
        v = ws.cell(row=r, column=c).value
        row_vals.append(v)
    print(f"Row {r}: {row_vals}")