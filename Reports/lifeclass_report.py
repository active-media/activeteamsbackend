#!/usr/bin/env python3
"""
generate_lifeclass_report.py
=============================
Generates the Life Class attendance report workbook from raw attendance data.

WHAT IT BUILDS
--------------
A 6-sheet workbook, in this order:
  1. MEN                    - raw attendance, weekly totals, Leader@12 stats table
  2. GUIDE STATE MEN         - weekly delegate count per guide, grouped by Leader@12
  3. GUIDE STATE FOR MEN     - weekly Guides / 1st-time / 2nd-time counts per Leader@12
  4. WOMEN                   - same as MEN, for the women's group
  5. GUIDE STATE WOMEN
  6. GUIDE STATE FOR WOMEN

Every derived number (weekly totals, leader stats, guide counts, delegate-type
breakdowns) is a live Excel formula (COUNTIF/COUNTIFS), so the report
recalculates automatically as attendance is marked in future weeks.

INPUT DATA FORMAT
------------------
Provide raw attendance data for MEN and WOMEN as CSV files (or as sheets in
an existing .xlsx - see --source-xlsx below). Each CSV/sheet needs these
columns, in this order:

    Name | Leader @12 | Leader @144 | Registration Type | WEEK 1 | WEEK 2 |
    WEEK 3 | WEEK 4 | ENCOUNTER | WEEK 5 | WEEK 6 | WEEK 7 | WEEK 8

  - Name:               the attendee's name (required; blank rows are skipped)
  - Leader @12:         the name of their Leader@12 (blank if the person IS a
                         top-level Leader@12 themselves)
  - Leader @144:        the name of their immediate guide, if different from
                         their Leader@12 (leave blank if their Leader@12 is
                         also their direct guide)
  - Registration Type:  any text containing "guide", "first", or "second"
                         (case-insensitive) - e.g. "Guide", "Delegate: First
                         Time", "Delegate: Second Time". Typos/spacing
                         variants are normalized automatically.
  - WEEK 1..WEEK 8, ENCOUNTER: mark attendance with "X" (any non-empty value
                         counts as present; leave blank if absent)

USAGE
-----
    # From two CSV files:
    python generate_lifeclass_report.py \\
        --men men_attendance.csv --women women_attendance.csv \\
        --output LifeClass_Report.xlsx

    # From an existing workbook that already has MEN / WOMEN sheets
    # (e.g. last term's report, or a raw export) in the column layout above:
    python generate_lifeclass_report.py \\
        --source-xlsx LastTerm.xlsx --men-sheet MEN --women-sheet WOMEN \\
        --output LifeClass_Report.xlsx

Requires: openpyxl (pip install openpyxl --break-system-packages)
After generating, recalculate the formulas with LibreOffice so the file
opens with values already computed - see the bottom of this file, or run:
    soffice --headless --convert-to xlsx --outdir . LifeClass_Report.xlsx
"""

import argparse
import csv
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WEEK_HEADERS = ['WEEK 1', 'WEEK 2', 'WEEK 3', 'WEEK 4', 'ENCOUNTER',
                'WEEK 5', 'WEEK 6', 'WEEK 7', 'WEEK 8']
N_WEEKS = len(WEEK_HEADERS)  # 9 week/encounter columns

FONT = 'Arial'


# --------------------------------------------------------------------------
# Data loading & normalization
# --------------------------------------------------------------------------

def norm_name(value):
    """Trim whitespace; return None for blank cells."""
    if value is None:
        return None
    value = str(value).strip()
    return value if value else None


def norm_regtype(value):
    """
    Collapse messy registration-type text into exactly one of:
      'Guide', 'Delegate: First Time', 'Delegate: Second Time'
    Matching is case-insensitive and tolerant of typos/spacing, e.g.
    'Delage: First Time ', 'Delegate:Second Time', 'Guide ' all normalize.
    """
    if value is None:
        return None
    v = str(value).lower()
    if 'guide' in v:
        return 'Guide'
    if 'first' in v:
        return 'Delegate: First Time'
    if 'second' in v:
        return 'Delegate: Second Time'
    return str(value).strip()


def read_rows_from_worksheet(ws):
    """Read attendee rows from an already-open openpyxl worksheet (data_only)."""
    rows = []
    for r in range(2, ws.max_row + 1):
        name = norm_name(ws.cell(r, 1).value)
        if not name:
            continue
        leader12 = norm_name(ws.cell(r, 2).value)
        leader144 = norm_name(ws.cell(r, 3).value)
        regtype = norm_regtype(ws.cell(r, 4).value)
        weeks = []
        for wc in range(5, 5 + N_WEEKS):
            v = ws.cell(r, wc).value
            weeks.append('X' if (v is not None and str(v).strip() != '') else None)
        rows.append({'name': name, 'leader12': leader12, 'leader144': leader144,
                      'regtype': regtype, 'weeks': weeks})
    return rows


def read_rows_from_csv(path):
    """Read attendee rows from a CSV file matching the documented column layout."""
    rows = []
    with open(path, newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        next(reader, None)  # skip header row
        for line in reader:
            if not line or not norm_name(line[0] if len(line) > 0 else None):
                continue
            line = line + [''] * (13 - len(line))  # pad short rows
            name = norm_name(line[0])
            leader12 = norm_name(line[1])
            leader144 = norm_name(line[2])
            regtype = norm_regtype(line[3])
            weeks = []
            for i in range(4, 4 + N_WEEKS):
                v = line[i].strip() if i < len(line) else ''
                weeks.append('X' if v else None)
            rows.append({'name': name, 'leader12': leader12, 'leader144': leader144,
                          'regtype': regtype, 'weeks': weeks})
    return rows


# --------------------------------------------------------------------------
# Styling helpers
# --------------------------------------------------------------------------

class Styles:
    def __init__(self):
        self.hdr_fill = PatternFill('solid', fgColor='1F4E78')
        self.hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        self.subhdr_fill = PatternFill('solid', fgColor='D9E1F2')
        self.subhdr_font = Font(name=FONT, bold=True, size=10)
        self.bold = Font(name=FONT, bold=True, size=10)
        self.normal = Font(name=FONT, size=10)
        thin = Side(style='thin', color='BFBFBF')
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.green_fill = PatternFill('solid', fgColor='C6EFCE')
        self.orange_fill = PatternFill('solid', fgColor='FCE4D6')
        self.red_fill = PatternFill('solid', fgColor='FFC7CE')
        self.center = Alignment(horizontal='center', vertical='center')
        self.reg_fill = {
            'Delegate: First Time': self.green_fill,
            'Delegate: Second Time': self.orange_fill,
            'Guide': self.red_fill,
        }


def style_header_row(ws, st, row, col_start, col_end):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row, c)
        cell.font = st.hdr_font
        cell.fill = st.hdr_fill
        cell.alignment = st.center
        cell.border = st.border


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --------------------------------------------------------------------------
# Sheet builders
# --------------------------------------------------------------------------

def build_attendance_sheet(wb, st, sheet_name, data, leader_label_a, leader_label_b):
    """
    Builds the raw attendance sheet (MEN or WOMEN):
      - one row per attendee with a live per-person TOTAL formula
      - a WEEKLY TOTAL row summing each week's attendance
      - a Leader @12 stats table tracking each leader's weekly headcount
      - a hidden 'Effective Guide' helper column (Leader@144, else Leader@12)
        that the Guide State sheets key off
    Returns (worksheet, first_data_row, last_data_row, ordered_list_of_leaders12)
    """
    ws = wb.create_sheet(sheet_name)
    headers = (['Name', leader_label_a, leader_label_b, 'Registration Type']
               + WEEK_HEADERS + ['TOTAL', 'Effective Guide'])
    for i, h in enumerate(headers, start=1):
        ws.cell(1, i, h)
    style_header_row(ws, st, 1, 1, len(headers))
    ws.freeze_panes = 'A2'

    first_data_row = 2
    for idx, rec in enumerate(data):
        r = first_data_row + idx
        ws.cell(r, 1, rec['name']).font = st.normal
        ws.cell(r, 2, rec['leader12']).font = st.normal
        ws.cell(r, 3, rec['leader144']).font = st.normal
        rc = ws.cell(r, 4, rec['regtype'])
        rc.font = st.normal
        if rec['regtype'] in st.reg_fill:
            rc.fill = st.reg_fill[rec['regtype']]
        for wi in range(N_WEEKS):
            cell = ws.cell(r, 5 + wi, rec['weeks'][wi])
            cell.alignment = st.center
            cell.font = st.normal

        week_start_col = get_column_letter(5)
        week_end_col = get_column_letter(4 + N_WEEKS)
        total_col = 5 + N_WEEKS  # column N
        ws.cell(r, total_col,
                f'=COUNTIF({week_start_col}{r}:{week_end_col}{r},"X")').font = st.bold

        eff_col = total_col + 1  # column O
        ws.cell(r, eff_col, f'=IF(C{r}="",B{r},C{r})').font = st.normal

        for c in range(1, eff_col + 1):
            ws.cell(r, c).border = st.border

    last_data_row = first_data_row + len(data) - 1

    # Weekly total row
    sum_row = last_data_row + 1
    ws.cell(sum_row, 4, 'WEEKLY TOTAL').font = st.bold
    ws.cell(sum_row, 4).alignment = Alignment(horizontal='right')
    for wi in range(N_WEEKS):
        col = 5 + wi
        cl = get_column_letter(col)
        cell = ws.cell(sum_row, col, f'=COUNTIF({cl}{first_data_row}:{cl}{last_data_row},"X")')
        cell.font = st.bold
        cell.fill = st.subhdr_fill
        cell.alignment = st.center
        cell.border = st.border
    total_col = 5 + N_WEEKS
    tcell = ws.cell(sum_row, total_col, f'=SUM(E{sum_row}:M{sum_row})')
    tcell.font = st.bold
    tcell.fill = st.subhdr_fill
    tcell.border = st.border

    # Leader @12 stats table
    leaders12 = list(OrderedDict.fromkeys(
        rec['leader12'] for rec in data if rec['leader12']))
    stats_title_row = sum_row + 3
    ws.cell(stats_title_row, 4, 'LEADER @12 STATS').font = st.bold
    hdr_row = stats_title_row + 1
    ws.cell(hdr_row, 4, 'Leader @12')
    for wi, wh in enumerate(WEEK_HEADERS):
        ws.cell(hdr_row, 5 + wi, wh)
    ws.cell(hdr_row, 5 + N_WEEKS, 'Total')
    style_header_row(ws, st, hdr_row, 4, 5 + N_WEEKS)

    for li, leader in enumerate(leaders12):
        r = hdr_row + 1 + li
        ws.cell(r, 4, leader).font = st.normal
        ws.cell(r, 4).border = st.border
        for wi in range(N_WEEKS):
            col = 5 + wi
            cl = get_column_letter(col)
            f = (f'=COUNTIFS($B${first_data_row}:$B${last_data_row},$D{r},'
                 f'{cl}${first_data_row}:{cl}${last_data_row},"X")')
            cell = ws.cell(r, col, f)
            cell.alignment = st.center
            cell.font = st.normal
            cell.border = st.border
        tcol = 5 + N_WEEKS
        tc = ws.cell(r, tcol, f'=SUM(E{r}:M{r})')
        tc.font = st.bold
        tc.border = st.border
        tc.alignment = st.center

    autosize(ws, [22, 20, 20, 22] + [9] * N_WEEKS + [9, 20])
    ws.column_dimensions['O'].hidden = True
    ws.sheet_view.showGridLines = False
    return ws, first_data_row, last_data_row, leaders12


def build_guide_state(wb, st, sheet_name, data, source_sheet_name, first_row, last_row, leaders12):
    """
    Weekly delegate count per guide, grouped by Leader@12. Uses the
    'Effective Guide' helper column on the source sheet so a leader who is
    also their own group's direct guide is counted correctly.
    """
    ws = wb.create_sheet(sheet_name)
    headers = ['Leader', 'Guide', 'Delegate W1', 'Delegate W2', 'Delegate W3',
               'Delegate W4', 'ENCOUNTER', 'Delegate W5', 'Delegate W6',
               'Delegate W7', 'Delegate W8']
    for i, h in enumerate(headers, start=1):
        ws.cell(1, i, h)
    style_header_row(ws, st, 1, 1, len(headers))
    ws.freeze_panes = 'A2'

    all_guide_names = {r['name'] for r in data if r['regtype'] == 'Guide'}
    guides_by_leader = OrderedDict()
    for leader in leaders12:
        glist = []
        if leader in all_guide_names:
            glist.append(leader)
        for rec in data:
            if (rec['regtype'] == 'Guide' and rec['leader12'] == leader
                    and rec['name'] not in glist):
                glist.append(rec['name'])
        guides_by_leader[leader] = glist

    src = f"'{source_sheet_name}'"
    r = 2
    for leader, glist in guides_by_leader.items():
        for gi, guide in enumerate(glist):
            ws.cell(r, 1, leader if gi == 0 else None).font = st.bold
            ws.cell(r, 2, guide).font = st.normal
            for wi in range(N_WEEKS):
                col = 3 + wi
                srccol = get_column_letter(5 + wi)
                f = (f'=COUNTIFS({src}!$O${first_row}:$O${last_row},$B{r},'
                     f'{src}!${srccol}${first_row}:${srccol}${last_row},"X")')
                cell = ws.cell(r, col, f)
                cell.alignment = st.center
                cell.font = st.normal
                cell.border = st.border
            ws.cell(r, 1).border = st.border
            ws.cell(r, 2).border = st.border
            r += 1

    autosize(ws, [22, 22] + [11] * N_WEEKS)
    ws.sheet_view.showGridLines = False
    return ws


def build_guide_state_for(wb, st, sheet_name, source_sheet_name, first_row, last_row, leaders12):
    """
    Per Leader@12, per week: counts of Guides / First-Time delegates /
    Second-Time delegates. Rendered as two stacked blocks (Weeks 1-4 +
    Encounter, then Weeks 5-8) to mirror the original report layout.
    """
    ws = wb.create_sheet(sheet_name)
    src = f"'{source_sheet_name}'"

    def write_block(top_row, week_indices):
        header_row = top_row
        sub_row = top_row + 1
        ws.cell(header_row, 1, 'Leader').font = st.bold
        col = 2
        for wi in week_indices:
            label = WEEK_HEADERS[wi] + ' DELEGATES'
            start_c = col
            ws.merge_cells(start_row=header_row, start_column=start_c,
                            end_row=header_row, end_column=start_c + 2)
            hc = ws.cell(header_row, start_c, label)
            hc.font = st.hdr_font
            hc.fill = st.hdr_fill
            hc.alignment = st.center
            for cc in range(start_c, start_c + 3):
                ws.cell(header_row, cc).fill = st.hdr_fill
                ws.cell(header_row, cc).border = st.border
            sub_labels = ['Guides', 'First Time', 'Second Time']
            for offset, lbl in enumerate(sub_labels):
                sc = ws.cell(sub_row, start_c + offset, lbl)
                sc.font = st.subhdr_font
                sc.fill = st.subhdr_fill
                sc.alignment = st.center
                sc.border = st.border
            col += 3
        ws.cell(sub_row, 1, 'Leader').font = st.subhdr_font
        ws.cell(sub_row, 1).fill = st.subhdr_fill
        ws.cell(sub_row, 1).border = st.border

        data_start = top_row + 2
        for li, leader in enumerate(leaders12):
            r = data_start + li
            ws.cell(r, 1, leader).font = st.normal
            ws.cell(r, 1).border = st.border
            col = 2
            for wi in week_indices:
                srccol = get_column_letter(5 + wi)
                base = (f'{src}!$B${first_row}:$B${last_row},$A{r},'
                        f'{src}!${srccol}${first_row}:${srccol}${last_row},"X",'
                        f'{src}!$D${first_row}:$D${last_row},')
                gcell = ws.cell(r, col, f'=COUNTIFS({base}"Guide")')
                fcell = ws.cell(r, col + 1, f'=COUNTIFS({base}"Delegate: First Time")')
                scell = ws.cell(r, col + 2, f'=COUNTIFS({base}"Delegate: Second Time")')
                for cell in (gcell, fcell, scell):
                    cell.font = st.normal
                    cell.alignment = st.center
                    cell.border = st.border
                col += 3
        return data_start + len(leaders12)

    end1 = write_block(1, [0, 1, 2, 3, 4])   # Weeks 1-4 + Encounter
    write_block(end1 + 2, [5, 6, 7, 8])      # Weeks 5-8

    autosize(ws, [22] + [10] * 15)
    ws.sheet_view.showGridLines = False
    return ws


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def build_report(men_rows, women_rows, output_path,
                  men_leader_labels=('Leader at 12', 'Leader at 144'),
                  women_leader_labels=('Leader 12', 'Leader at 144')):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    st = Styles()

    men_ws, men_first, men_last, men_leaders = build_attendance_sheet(
        wb, st, 'MEN', men_rows, *men_leader_labels)
    build_guide_state(wb, st, 'GUIDE STATE MEN', men_rows, 'MEN',
                       men_first, men_last, men_leaders)
    build_guide_state_for(wb, st, 'GUIDE STATE FOR MEN', 'MEN',
                           men_first, men_last, men_leaders)

    women_ws, women_first, women_last, women_leaders = build_attendance_sheet(
        wb, st, 'WOMEN', women_rows, *women_leader_labels)
    build_guide_state(wb, st, 'GUIDE STATE WOMEN', women_rows, 'WOMEN',
                       women_first, women_last, women_leaders)
    build_guide_state_for(wb, st, 'GUIDE STATE FOR WOMEN', 'WOMEN',
                           women_first, women_last, women_leaders)

    wb.active = 0
    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(
        description='Generate the Life Class attendance report workbook.',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__)
    parser.add_argument('--men', help="CSV file with the men's raw attendance data")
    parser.add_argument('--women', help="CSV file with the women's raw attendance data")
    parser.add_argument('--source-xlsx',
                         help='Existing .xlsx to read MEN/WOMEN raw data from, instead of CSVs')
    parser.add_argument('--men-sheet', default='MEN',
                         help="Sheet name for the men's data in --source-xlsx (default: MEN)")
    parser.add_argument('--women-sheet', default='WOMEN',
                         help="Sheet name for the women's data in --source-xlsx (default: WOMEN)")
    parser.add_argument('--output', default='LifeClass_Report.xlsx',
                         help='Output workbook path (default: LifeClass_Report.xlsx)')
    args = parser.parse_args()

    if args.source_xlsx:
        src_wb = openpyxl.load_workbook(args.source_xlsx, data_only=True)
        men_rows = read_rows_from_worksheet(src_wb[args.men_sheet])
        women_rows = read_rows_from_worksheet(src_wb[args.women_sheet])
    elif args.men and args.women:
        men_rows = read_rows_from_csv(args.men)
        women_rows = read_rows_from_csv(args.women)
    else:
        parser.error('Provide either --source-xlsx, or both --men and --women CSV files.')
        return

    out = build_report(men_rows, women_rows, args.output)
    print(f'Report written to: {out}')
    print(f'  MEN:   {len(men_rows)} attendees')
    print(f'  WOMEN: {len(women_rows)} attendees')
    print('Formulas are written but uncalculated - open in Excel/Google Sheets '
          '(they recalc on open), or run LibreOffice headless to bake in values:')
    print(f'  soffice --headless --convert-to xlsx --outdir . {out}')


if __name__ == '__main__':
    main()