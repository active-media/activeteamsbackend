from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timedelta, timezone
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from supabase_helpers.supabase_connection import supabase

# ---------------------------------------------------------------------------
# NOTE / ASSUMPTION: this module assumes Tasks.assignedfor stores the
# assignee's People/Users _id (a stable id), not a display name or raw
# email string. The original code grouped by a lowercased name string,
# which is fragile (casing/typo mismatches create phantom "people"). If
# assignedfor is actually storing email, swap PERSON_KEY_FIELDS below to
# match on email instead, but a stable id is strongly preferred since the
# roster table keys on person_id.
# ---------------------------------------------------------------------------
PERSON_KEY_FIELDS = ("assignedfor", "assigned_to_email")

# Indicator colors — must match the legend in the exported report.
FILL_ON_TARGET = PatternFill("solid", fgColor="FFC000")     # orange
FILL_LESS_THAN_TARGET = PatternFill("solid", fgColor="FF0000")  # red
FILL_MORE_THAN_TARGET = PatternFill("solid", fgColor="00B050")  # green

HEADER_FONT = Font(bold=True)


def _period_range(period: str) -> tuple[datetime, datetime]:
    """Return (start, end) UTC-aware datetimes for the requested period."""
    now = datetime.now(timezone.utc)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)

    if period in ("today", "daily"):
        return today, today.replace(hour=23, minute=59, second=59, microsecond=999_999)

    if period in ("thisWeek", "weekly"):
        start = today - timedelta(days=today.weekday())  # Monday
        return start, start + timedelta(days=6, hours=23, minutes=59, seconds=59)

    if period in ("thisMonth", "monthly"):
        start = today.replace(day=1)
        if today.month == 12:
            end = datetime(today.year + 1, 1, 1, tzinfo=timezone.utc) - timedelta(microseconds=1)
        else:
            end = datetime(today.year, today.month + 1, 1, tzinfo=timezone.utc) - timedelta(microseconds=1)
        return start, end

    if period == "previousWeek":
        last = today - timedelta(weeks=1)
        start = last - timedelta(days=last.weekday())
        return start, start + timedelta(days=6, hours=23, minutes=59, seconds=59)

    if period == "previousMonth":
        year, month = today.year, today.month - 1
        if month == 0:
            year, month = year - 1, 12
        start = datetime(year, month, 1, tzinfo=timezone.utc)
        if month == 12:
            end = datetime(year + 1, 1, 1, tzinfo=timezone.utc) - timedelta(microseconds=1)
        else:
            end = datetime(year, month + 1, 1, tzinfo=timezone.utc) - timedelta(microseconds=1)
        return start, end

    raise ValueError(f"Unknown period: {period!r}")


def _iso(dt: datetime) -> str:
    """Return an ISO-8601 string with timezone offset (Supabase-compatible)."""
    return dt.isoformat()


# ---------------------------------------------------------------------------
# Roster (Column A) — persisted, admin-managed list of tracked people.
# Independent of live LeaderPath/hierarchy so it only changes when an admin
# explicitly adds/removes someone. Removing someone stops future tracking
# but never deletes their historical rows (soft delete via `active`).
# ---------------------------------------------------------------------------

def sb_get_tracked_people(leader12_id: Optional[str] = None, org_filter: Optional[dict] = None) -> list[dict]:
    """Return the persisted, admin-curated roster for the Twelve Tasks report.

    Only rows with active=True are returned — this is what should populate
    Column A. Historical report rows for removed people are preserved
    elsewhere (task history), just not re-included in future roster pulls.
    """
    query = (
        supabase.table("twelve_tasks_tracked_people")
        .select("person_id, leader12_id, org, display_name, added_at")
        .eq("active", True)
    )
    if leader12_id:
        query = query.eq("leader12_id", leader12_id)

    rows = query.execute().data or []

    # Supabase SDK rule: never filter on joined-table columns with .eq()/.ilike();
    # fetch and filter Organization in Python instead.
    if org_filter:
        org_value = org_filter.get("Organization") or org_filter.get("organization")
        if org_value:
            rows = [r for r in rows if r.get("org") == org_value]

    return rows


def sb_add_tracked_person(person_id: str, leader12_id: str, display_name: str, org: str) -> dict:
    """Admin action: add a person to the tracked roster. Starts at 0/0 —
    caller does not need to seed history rows."""
    existing = (
        supabase.table("twelve_tasks_tracked_people")
        .select("person_id")
        .eq("person_id", person_id)
        .eq("leader12_id", leader12_id)
        .execute()
        .data
    )
    if existing:
        # Re-activating someone previously removed — never insert a duplicate row.
        return (
            supabase.table("twelve_tasks_tracked_people")
            .update({"active": True, "display_name": display_name, "org": org})
            .eq("person_id", person_id)
            .eq("leader12_id", leader12_id)
            .execute()
            .data
        )

    return (
        supabase.table("twelve_tasks_tracked_people")
        .insert(
            {
                "person_id": person_id,
                "leader12_id": leader12_id,
                "display_name": display_name,
                "org": org,
                "active": True,
                "added_at": _iso(datetime.now(timezone.utc)),
            }
        )
        .execute()
        .data
    )


def sb_remove_tracked_person(person_id: str, leader12_id: str) -> dict:
    """Admin action: soft-remove a person. Stops future tracking only —
    their B/D history stays intact in past report snapshots/exports."""
    return (
        supabase.table("twelve_tasks_tracked_people")
        .update({"active": False})
        .eq("person_id", person_id)
        .eq("leader12_id", leader12_id)
        .execute()
        .data
    )


# ---------------------------------------------------------------------------
# Targets (Column C) — per-person, admin/leader-editable, persists until
# explicitly changed. Independent of any other report's target field.
# ---------------------------------------------------------------------------

def sb_get_targets(person_ids: list[str]) -> dict[str, int]:
    if not person_ids:
        return {}
    rows = (
        supabase.table("twelve_tasks_targets")
        .select("person_id, target")
        .in_("person_id", person_ids)
        .execute()
        .data
        or []
    )
    return {row["person_id"]: row["target"] for row in rows}


def sb_set_target(person_id: str, target: int, updated_by: Optional[str] = None) -> dict:
    """Admin/leader action: set or change a person's target. Persists as-is
    until this is called again — never reset automatically."""
    payload = {
        "person_id": person_id,
        "target": target,
        "updated_at": _iso(datetime.now(timezone.utc)),
    }
    if updated_by:
        payload["updated_by"] = updated_by

    existing = (
        supabase.table("twelve_tasks_targets").select("person_id").eq("person_id", person_id).execute().data
    )
    if existing:
        return (
            supabase.table("twelve_tasks_targets")
            .update(payload)
            .eq("person_id", person_id)
            .execute()
            .data
        )
    return supabase.table("twelve_tasks_targets").insert(payload).execute().data


# ---------------------------------------------------------------------------
# Task counting — fixed to only count actually-COMPLETED tasks, using
# completedAt as the sole date field. Previously this also matched on
# followup_date (due date) and created_at, which inflated counts with
# tasks that were merely due or newly created, not completed.
# ---------------------------------------------------------------------------

def _count_completed_by_person(tasks: list[dict], start: datetime, end: datetime) -> dict[str, int]:
    start_iso, end_iso = _iso(start), _iso(end)
    counts: dict[str, int] = defaultdict(int)

    for task in tasks:
        if task.get("status") != "completed":
            continue

        completed_at = str(task.get("completedAt") or "")
        if not (start_iso <= completed_at <= end_iso):
            continue

        person_key = None
        for field in PERSON_KEY_FIELDS:
            value = (task.get(field) or "").strip()
            if value:
                person_key = value
                break
        if person_key:
            counts[person_key] += 1

    return dict(counts)


def _fetch_completed_tasks(start: datetime, end: datetime, org_filter: Optional[dict] = None) -> list[dict]:
    task_q = (
        supabase.table("Tasks")
        .select("_id, name, taskType, status, completedAt, assignedfor, assigned_to_email")
        .eq("status", "completed")
        .gte("completedAt", _iso(start))
        .lte("completedAt", _iso(end))
    )

    if org_filter:
        org_value = org_filter.get("Organization") or org_filter.get("organization")
        if org_value:
            task_q = task_q.eq("Organization", org_value)

    return task_q.execute().data or []


# ---------------------------------------------------------------------------
# Report assembly — roster-driven (left join), so everyone on the tracked
# list appears even with 0 completed tasks, and the row order follows the
# roster rather than being re-sorted by count.
# ---------------------------------------------------------------------------

def sb_get_twelve_tasks_report(leader12_id: Optional[str] = None, org_filter: Optional[dict] = None) -> dict:
    """
    Build the Twelve Tasks report:
      - Column A (name)        -> persisted roster (sb_get_tracked_people)
      - Column B (last week)   -> completed-task count, previous Mon-Sun
      - Column C (targets)     -> persisted per-person target
      - Column D (this week)   -> completed-task count, current Mon-Sun (so far)
    """
    this_week_start, this_week_end = _period_range("thisWeek")
    last_week_start, last_week_end = _period_range("previousWeek")

    roster = sb_get_tracked_people(leader12_id=leader12_id, org_filter=org_filter)
    person_ids = [r["person_id"] for r in roster]
    targets = sb_get_targets(person_ids)

    # Fetch once across the full span, then split in Python (cheaper than two round trips).
    tasks = _fetch_completed_tasks(last_week_start, this_week_end, org_filter=org_filter)
    last_week_counts = _count_completed_by_person(tasks, last_week_start, last_week_end)
    this_week_counts = _count_completed_by_person(tasks, this_week_start, this_week_end)

    rows = []
    for person in roster:
        pid = person["person_id"]
        rows.append(
            {
                "person_id": pid,
                "name": person.get("display_name") or pid,
                "last_week": last_week_counts.get(pid, 0),
                "target": targets.get(pid, 0),
                "this_week": this_week_counts.get(pid, 0),
            }
        )

    return {
        "period": {"start": this_week_start.date().isoformat(), "end": this_week_end.date().isoformat()},
        "previous_period": {
            "start": last_week_start.date().isoformat(),
            "end": last_week_end.date().isoformat(),
        },
        "rows": rows,
    }


# ---------------------------------------------------------------------------
# XLSX export — reproduces the exact layout from the mockup: Name / Last
# week / Targets / Actual this week, plus the Key legend, plus the
# orange/red/green indicator coloring on both the "Last week" and
# "Actual this week" columns (compared against that row's target).
# ---------------------------------------------------------------------------

def _indicator_fill(value: int, target: int) -> Optional[PatternFill]:
    if target == 0:
        return None  # no target set yet — leave uncolored rather than guessing
    if value == target:
        return FILL_ON_TARGET
    if value < target:
        return FILL_LESS_THAN_TARGET
    return FILL_MORE_THAN_TARGET


def export_twelve_tasks_xlsx(report: dict, path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = ["Name", " Last week ", "Targets", "Actual  this week"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT

    for row_idx, row in enumerate(report["rows"], start=2):
        ws.cell(row=row_idx, column=1, value=row["name"])

        last_week_cell = ws.cell(row=row_idx, column=2, value=row["last_week"])
        ws.cell(row=row_idx, column=3, value=row["target"])
        this_week_cell = ws.cell(row=row_idx, column=4, value=row["this_week"])

        fill = _indicator_fill(row["last_week"], row["target"])
        if fill:
            last_week_cell.fill = fill

        fill = _indicator_fill(row["this_week"], row["target"])
        if fill:
            this_week_cell.fill = fill

    # Key / legend block, matching the mockup (columns F/G)
    ws.cell(row=1, column=6, value="Key")
    legend = [
        (FILL_ON_TARGET, "On target"),
        (FILL_LESS_THAN_TARGET, "Less than target"),
        (FILL_MORE_THAN_TARGET, "More than target"),
    ]
    for i, (fill, label) in enumerate(legend, start=2):
        ws.cell(row=i, column=6).fill = fill
        ws.cell(row=i, column=7, value=label)

    for col, width in zip("ABCDFG", (16, 12, 10, 16, 6, 20)):
        ws.column_dimensions[col].width = width

    wb.save(path)
    return path