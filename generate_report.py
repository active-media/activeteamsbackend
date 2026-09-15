import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime, timedelta
from supabase_helpers.twelve_tasks import sb_get_twelve_tasks_report, _period_range, _iso
from unittest.mock import MagicMock, patch


def generate_excel_report(period="thisWeek", output_path="twelve_tasks_report.xlsx"):
    """Generate the Twelve Tasks report and export to Excel."""

    # Generate the report with mocked Supabase
    with patch("supabase_helpers.twelve_tasks.supabase") as mock_supabase:
        start, end = _period_range(period)
        print(f"DEBUG: period={period}, start={start}, end={end}")

        current_tasks = [
            {"assignedfor": "leader1@example.com", "followup_date": _iso(start)},
            {"assignedfor": "leader1@example.com", "followup_date": _iso(start + timedelta(days=5))},
            {"assignedfor": "leader2@example.com", "followup_date": _iso(start + timedelta(days=3))},
            {"assignedfor": "leader3@example.com", "followup_date": _iso(start + timedelta(days=1))},
        ]
        previous_tasks = [
            {"assignedfor": "leader1@example.com", "followup_date": _iso(start - timedelta(days=1))},
            {"assignedfor": "leader2@example.com", "followup_date": _iso(start - timedelta(days=4))},
            {"assignedfor": "leader3@example.com", "followup,
        ]

        mock_cursor_current = MagicMock()
        mock_cursor_current.execute.return_value.data = current_tasks

        mock_cursor_previous = MagicMock()
        mock_cursor_previous.execute.return_value.data = previous_tasks

        with patch.object(
            mock_supabase.table.return_value.select.return_value.or_,
            "execute",
            side_effect=[mock_cursor_current, mock_cursor_previous],
        ):
            result = sb_get_twelve_tasks_report(period=period)
            print(f"DEBUG: result={result}")
            print(f"DEBUG: leaders={result.get('leaders', [])}")

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Twelve Tasks Report"

    # Style definitions
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    headers = ["Leader Name", "Total Captured", "Previous Total", "Change", "Change %"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write leader data
    leaders = result.get("leaders", [])
    print(f"DEBUG: About to write {len(leaders)} leaders")
    for i, leader in enumerate(leaders, 2):
        print(f"DEBUG: Writing leader {leader}")
        ws.cell(row=i, column=1, value=leader.get("name", "")).border = thin_border
        ws.cell(row=i, column=2, value=leader.get("total_captured", 0)).border = thin_border
        ws.cell(row=i, column=3, value=leader.get("previous_total_captured", 0)).border = thin_border
        ws.cell(row=i, column=4, value=leader.get("change", 0)).border = thin_border
        change_pct = leader.get("change_percent")
        ws.cell(
            row=i, column=5, value=change_pct if change_pct is not None else "N/A"
        ).border = thin_border

    # Adjust column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 15

    # Write summary section
    ws["A8"] = "Report Summary"
    ws["A8"].font = Font(bold=True, size=13)

    period_data = result.get("period", {})
    prev_data = result.get("previous_period", {})
    ws["A9"] = f"Period: {period_data.get('start', 'N/A')} to {period_data.get('end', 'N/A')}"
    ws["A10"] = f"Previous Period: {prev_data.get('start', 'N/A')} to {prev_data.get('end', 'N/A')}"
    ws["A11"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    # Save Excel file
    wb.save(output_path)
    print(f"Excel report saved to: {output_path}")

    # Print summary
    print(f"\nLeaders ({len(leaders)}):")
    for leader in leaders:
        print(
            f"  {leader['name']}: {leader['total_captured']} current, "
            f"{leader['previous_total_captured']} previous, "
            f"Change: {leader['change']}, %: {leader['change_percent']}"
        )

    return result, output_path


if __name__ == "__main__":
    generate_excel_report()